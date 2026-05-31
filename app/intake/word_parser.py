from io import BytesIO


def parse_docx_bytes(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to parse Word files.") from exc

    document = Document(BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_cells: list[str] = []

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    table_cells.append(cell.text.strip())

    return "\n".join(paragraphs + table_cells).strip()

